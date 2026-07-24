"""Tests for resume PDF/DOCX markdown export parsing and rendering."""

from django.test import SimpleTestCase

from resume_app.api import _build_export_docx, _build_export_pdf, _parse_markdown_blocks


class ParseMarkdownBlocksTests(SimpleTestCase):
    def test_parses_heading_levels_one_through_six(self):
        content = "\n".join(
            [
                "# Name",
                "## EXPERIENCE",
                "### Senior Engineer | Acme",
                "#### Team Lead",
                "##### Nested",
                "###### Deep",
                "- Built systems",
                "Summary paragraph",
            ]
        )
        blocks = list(_parse_markdown_blocks(content))
        self.assertEqual(
            blocks,
            [
                ("heading1", "Name"),
                ("heading2", "EXPERIENCE"),
                ("heading3", "Senior Engineer | Acme"),
                ("heading4", "Team Lead"),
                ("heading5", "Nested"),
                ("heading6", "Deep"),
                ("bullet", "Built systems"),
                ("paragraph", "Summary paragraph"),
            ],
        )

    def test_hash_markers_are_stripped_from_heading_text(self):
        blocks = list(_parse_markdown_blocks("### Role Title\n#### Subrole"))
        self.assertEqual(blocks[0], ("heading3", "Role Title"))
        self.assertEqual(blocks[1], ("heading4", "Subrole"))
        self.assertNotIn("###", blocks[0][1])
        self.assertNotIn("####", blocks[1][1])


class ExportBuildersTests(SimpleTestCase):
    def test_pdf_export_accepts_h3_h4_without_literal_hashes(self):
        content = "## EXPERIENCE\n### Senior Engineer\n#### Platform Team\n- Shipped features"
        buf = _build_export_pdf(content)
        self.assertIsNotNone(buf)
        pdf_bytes = buf.getvalue()
        self.assertTrue(pdf_bytes.startswith(b"%PDF"))
        # Raw markdown markers should not appear in the PDF stream as plain text.
        self.assertNotIn(b"### Senior", pdf_bytes)
        self.assertNotIn(b"#### Platform", pdf_bytes)

    def test_docx_export_maps_h3_h4_to_heading_styles(self):
        from docx import Document

        content = "## EXPERIENCE\n### Senior Engineer\n#### Platform Team\n- Shipped features"
        buf = _build_export_docx(content)
        self.assertIsNotNone(buf)
        doc = Document(buf)
        styles = [p.style.name for p in doc.paragraphs]
        texts = [p.text for p in doc.paragraphs]
        self.assertIn("Heading 1", styles)  # ## -> level 1
        self.assertIn("Heading 2", styles)  # ### -> level 2
        self.assertIn("Heading 3", styles)  # #### -> level 3
        self.assertEqual(texts[0], "EXPERIENCE")
        self.assertEqual(texts[1], "Senior Engineer")
        self.assertEqual(texts[2], "Platform Team")
        self.assertTrue(all("###" not in t and "####" not in t for t in texts))
