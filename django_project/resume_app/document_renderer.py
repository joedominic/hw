"""
Document Rendering Service.
Compiles markdown-style formatted text into downloadable PDF (ReportLab) and Word (DOCX) files.
"""
import io
import re
from typing import TYPE_CHECKING
from django.http import FileResponse

if TYPE_CHECKING:
    from .models import OptimizedResume


def parse_markdown_blocks(content: str):
    """Yield (block_type, text).

    block_type: heading1..heading6, bullet, paragraph.
    Heading markers must be checked longest-first so ### / #### are not left as plain text.
    """
    if not content:
        return
    heading_re = re.compile(r"^(#{1,6})\s+(.+)$")
    for line in (content or "").replace("\r", "").split("\n"):
        raw = line
        line = line.strip()
        if not line:
            continue
        heading_match = heading_re.match(line)
        if heading_match:
            level = len(heading_match.group(1))
            yield (f"heading{level}", heading_match.group(2).strip())
        elif line.startswith("- ") or line.startswith("* "):
            yield ("bullet", line[2:].strip())
        else:
            yield ("paragraph", raw.strip())


def normalize_export_content(text: str) -> str:
    """Normalize common Unicode punctuation and whitespace for PDF/Word export."""
    if not text:
        return text or ""
    from .utils import sanitize_resume_markdown
    return sanitize_resume_markdown(text)


def split_style_spans(text: str):
    """Yield (is_bold, is_italic, segment) for text with **bold** and *italic* markers."""
    parts = re.split(r"(\*\*\*[^*]+\*\*\*|\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for p in parts:
        if not p:
            continue
        if p.startswith("***") and p.endswith("***"):
            yield (True, True, p[3:-3])
        elif p.startswith("**") and p.endswith("**"):
            yield (True, False, p[2:-2])
        elif p.startswith("*") and p.endswith("*"):
            yield (False, True, p[1:-1])
        else:
            yield (False, False, p)


def draw_rich_text(
    c,
    x: float,
    y: float,
    max_width: float,
    segments,
    line_height: float,
    font_size: float = 10,
    first_line_indent: float = None,
    default_font_name: str = "Times-Roman",
    text_color = None,
) -> float:
    if line_height is None:
        line_height = font_size * 1.5
    if first_line_indent is None:
        first_line_indent = x
    current_x = x
    current_y = y
    line_start = x
    for is_bold, is_italic, segment in segments:
        tokens = re.findall(r"\s+|\S+", segment)
        for token in tokens:
            if is_bold and is_italic:
                if "Times" in default_font_name:
                    font_name = "Times-BoldItalic"
                else:
                    font_name = "Helvetica-BoldOblique"
            elif is_bold:
                if "Times" in default_font_name:
                    font_name = "Times-Bold"
                else:
                    font_name = "Helvetica-Bold"
            elif is_italic:
                if "Times" in default_font_name:
                    font_name = "Times-Italic"
                else:
                    font_name = "Helvetica-Oblique"
            else:
                font_name = default_font_name
            c.setFont(font_name, font_size)
            if text_color:
                c.setFillColor(text_color)
            token_width = c.stringWidth(token, font_name, font_size)
            if token.strip() and current_x + token_width > line_start + max_width:
                current_y -= line_height
                current_x = first_line_indent
                line_start = first_line_indent
            if token.isspace() and current_x == line_start:
                continue
            c.drawString(current_x, current_y, token)
            current_x += token_width
    from reportlab.lib import colors
    c.setFillColor(colors.black)
    return current_y - line_height


def build_export_pdf(content: str) -> io.BytesIO | None:
    """Build PDF from markdown-style content with headings, bullets, bold, italic.
    Uses Times-Roman for professional typography and blue headings for visual hierarchy."""
    content = normalize_export_content(content)
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas
        from reportlab.lib.units import inch
        from reportlab.lib import colors
        
        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=letter)
        width, height = letter
        y = height - inch
        body_line_height = 15
        text_width = width - 2 * inch
        heading_color = (31/255, 78/255, 121/255)
        
        heading_styles = {
            "heading1": {"font_size": 16, "block_height": 18, "space_before": 14, "space_after": 2},
            "heading2": {"font_size": 13, "block_height": 15, "space_before": 9, "space_after": 2},
            "heading3": {"font_size": 12, "block_height": 14, "space_before": 6, "space_after": 2},
            "heading4": {"font_size": 10, "block_height": 12, "space_before": 4, "space_after": 2},
            "heading5": {"font_size": 10, "block_height": 12, "space_before": 4, "space_after": 2},
            "heading6": {"font_size": 10, "block_height": 12, "space_before": 4, "space_after": 2},
        }

        for block_type, text in parse_markdown_blocks(content):
            heading_style = heading_styles.get(block_type)
            block_height = body_line_height
            if heading_style:
                block_height = heading_style["block_height"]
            elif block_type == "bullet":
                block_height = 16

            if y < inch + block_height * 3:
                c.showPage()
                y = height - inch

            if heading_style:
                if y < height - inch - 2:
                    y -= heading_style.get("space_before", 6)
                y = draw_rich_text(
                    c,
                    inch,
                    y,
                    text_width,
                    split_style_spans(text),
                    block_height,
                    font_size=heading_style["font_size"],
                    default_font_name="Times-Bold",
                    text_color=heading_color,
                )
                y -= heading_style["space_after"]
            elif block_type == "bullet":
                c.setFont("Times-Roman", 10)
                c.setFillColor(colors.black)
                c.drawString(inch, y, "•")
                y = draw_rich_text(
                    c,
                    inch + 14,
                    y,
                    text_width - 14,
                    split_style_spans(text),
                    block_height,
                    font_size=10,
                    first_line_indent=inch + 14,
                    default_font_name="Times-Roman",
                )
                y -= body_line_height / 6
            else:
                y = draw_rich_text(
                    c,
                    inch,
                    y,
                    text_width,
                    split_style_spans(text),
                    body_line_height,
                    font_size=10,
                    default_font_name="Times-Roman",
                )
                y -= body_line_height / 5

        c.save()
        buf.seek(0)
        return buf
    except ImportError:
        return None


def build_export_docx(content: str) -> io.BytesIO | None:
    """Build Word from markdown-style content with headings, bullets, bold, italic."""
    content = normalize_export_content(content)
    try:
        from docx import Document
        from docx.shared import Pt
        doc = Document()
        heading_levels = {
            "heading1": 0,
            "heading2": 1,
            "heading3": 2,
            "heading4": 3,
            "heading5": 4,
            "heading6": 5,
        }
        for block_type, text in parse_markdown_blocks(content):
            if block_type in heading_levels:
                p = doc.add_heading(text, level=heading_levels[block_type])
                p.paragraph_format.space_after = Pt(8 if block_type == "heading1" else 6)
            elif block_type == "bullet":
                p = doc.add_paragraph(style="List Bullet")
                for is_bold, is_italic, segment in split_style_spans(text):
                    r = p.add_run(segment)
                    r.bold = is_bold
                    r.italic = is_italic
                p.paragraph_format.space_after = Pt(4)
            else:
                p = doc.add_paragraph()
                for is_bold, is_italic, segment in split_style_spans(text):
                    r = p.add_run(segment)
                    r.bold = is_bold
                    r.italic = is_italic
                p.paragraph_format.space_after = Pt(6)
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf
    except ImportError:
        return None


def apply_export_replacements(content: str, request) -> str:
    if not content:
        return content
    replacements = []
    params = request.GET
    for key, value in params.items():
        if not key.startswith("replace_token_") or not value:
            continue
        suffix = key[len("replace_token_"):]
        replacement = params.get(f"replace_value_{suffix}")
        if replacement is None or replacement == "":
            continue
        replacements.append((value, replacement))

    if not replacements:
        from .models import AppAutomationSettings

        raw_replacements = AppAutomationSettings.get_for_user(request.user).export_replacements or []
        if not raw_replacements:
            raw_replacements = request.session.get("export_replacements") or []
        for entry in raw_replacements:
            if not isinstance(entry, dict):
                continue
            token = (entry.get("token") or "").strip()
            replacement = entry.get("value")
            if not token or replacement is None or replacement == "":
                continue
            replacements.append((token, replacement))

    for token, replacement in replacements:
        content = content.replace(token, replacement)
    return content


def export_file_response(buf: io.BytesIO, *, filename: str, content_type: str) -> FileResponse:
    """Return a download response that must not be served from browser/proxy cache."""
    response = FileResponse(buf, as_attachment=True, filename=filename, content_type=content_type)
    response["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response["Pragma"] = "no-cache"
    response["Expires"] = "0"
    return response


def format_export_filename(prefix: str, company: str, title: str, ext: str) -> str:
    parts = [prefix]
    if company:
        parts.append(company)
    if title:
        parts.append(title)
    base = "_".join(parts)
    cleaned = re.sub(r"[^\w\s-]", "", base).strip()
    cleaned = re.sub(r"[-\s]+", "_", cleaned)
    return f"{cleaned or prefix}.{ext}"


def job_meta_from_optimized_safe(optimized: "OptimizedResume") -> tuple[str, str, str]:
    try:
        entry = getattr(optimized, "pipeline_entry", None)
        if entry is not None and getattr(entry, "job_listing", None) is not None:
            job = entry.job_listing
            return (
                (job.company_name or "").strip(),
                (job.title or "").strip(),
                (job.url or "").strip(),
            )
    except Exception:
        pass
    return ("", "", "")


# Backwards-compatible aliases
_parse_markdown_blocks = parse_markdown_blocks
_normalize_export_content = normalize_export_content
_split_style_spans = split_style_spans
_draw_rich_text = draw_rich_text
_build_export_pdf = build_export_pdf
_build_export_docx = build_export_docx
_apply_export_replacements = apply_export_replacements
_export_file_response = export_file_response
_format_export_filename = format_export_filename
_job_meta_from_optimized_safe = job_meta_from_optimized_safe
